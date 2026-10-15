"""Offline regressions for manifest-bound, loss-preserving data import."""
from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

REPO = Path(__file__).resolve().parents[1]
SCRIPTS = REPO / 'packages/claude/.claude/skills/data-cleaning-and-visualization/scripts'


class DataIntegrityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='mathmodel-data-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        (self.root / 'problem_files').mkdir()
        (self.root / 'paper_output').mkdir()

    def source(self, name='data.csv', text='id,value\n001,1\n002,\n003,9\n'):
        path = self.root / 'problem_files' / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding='utf-8')
        return path

    def manifest(self, files, excluded=()):
        entries = [{'path': p.relative_to(self.root).as_posix(),
                    'role': 'result_template' if p in excluded else 'raw_data',
                    'usable_for_modeling': p not in excluded,
                    'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in files]
        path = self.root / 'paper_output/input_manifest.json'
        path.write_text(json.dumps({'entries': entries}), encoding='utf-8')
        return path

    def run_script(self, name='clean_data.py', *args):
        return subprocess.run([sys.executable, str(SCRIPTS / name), *args], cwd=self.root,
                              env={**os.environ, 'PYTHONDONTWRITEBYTECODE': '1', 'PYTHONUTF8': '1'},
                              capture_output=True, text=True, encoding='utf-8', timeout=45)

    def report(self):
        return json.loads((self.root / 'paper_output/data_cleaned/cleaning_report.json').read_text(encoding='utf-8'))

    def imported_rows(self, dataset):
        with (self.root / dataset['cleaned_output']).open(encoding='utf-8-sig', newline='') as handle:
            return list(csv.reader(handle))

    def test_templates_and_unlisted_files_never_imported(self):
        data = self.source()
        excluded = self.source('answer.csv', 'result\n999\n')
        self.source('unlisted.csv', 'secret\n888\n')
        self.manifest([data, excluded], [excluded])
        result = self.run_script()
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        report = self.report()
        self.assertEqual([d['path'] for d in report['datasets']], ['problem_files/data.csv'])
        self.assertEqual(report['skipped_files'][0]['role'], 'result_template')

    def test_values_missing_rows_identifiers_and_duplicates_preserved(self):
        data = self.source(text='id,value,tag\n001,1,NA\n002,,null\n003,9,x\n003,9,x\n\n')
        self.manifest([data])
        result = self.run_script()
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        item = self.report()['datasets'][0]
        self.assertEqual(self.imported_rows(item), [
            ['id', 'value', 'tag'], ['001', '1', 'NA'], ['002', '', 'null'],
            ['003', '9', 'x'], ['003', '9', 'x'], ['', '', '']])
        self.assertEqual(item['rows_before'], 5)
        self.assertEqual(item['rows_after'], 5)
        self.assertEqual(item['duplicate_rows_preserved'], 1)
        self.assertEqual(item['transformations'], [])

    def test_bad_width_and_broken_quotes_fail_without_silent_skip(self):
        for text in ('a,b\n1,2\n3,4,5\n', 'a,b\n1\n', 'a,b\n1,"unterminated'):
            with self.subTest(text=text):
                self.manifest([self.source(text=text)])
                result = self.run_script()
                self.assertNotEqual(result.returncode, 0)
                self.assertEqual(self.report()['status'], 'FAIL')
                self.assertEqual(self.report()['datasets'], [])

    def test_delimiters_encoding_and_quoted_newlines(self):
        sys.path.insert(0, str(SCRIPTS))
        from clean_data import read_tables
        for name, text, encoding, expected in (
            ('semi.csv', 'id;value\n001;2\n', 'utf-8', ['001', '2']),
            ('tab.tsv', 'id\tvalue\n001\t2\n', 'utf-8', ['001', '2']),
            ('quote.csv', 'id,note\n001,"a,b\nnext line"\n', 'utf-8', ['001', 'a,b\nnext line']),
            ('中文.csv', '编号,名称\n001,样本\n', 'gb18030', ['001', '样本']),
        ):
            with self.subTest(name=name):
                path = self.root / 'problem_files' / name
                path.write_bytes(text.encode(encoding))
                _, frame, _ = read_tables(path)[0]
                self.assertEqual(frame.iloc[0].tolist(), expected)

    def test_every_excel_sheet_is_recorded(self):
        path = self.root / 'problem_files/workbook.xlsx'
        wb = Workbook()
        for index, name in enumerate(('训练', '验证')):
            sheet = wb.active if index == 0 else wb.create_sheet()
            sheet.title = name
            sheet.append(['id', 'value'])
            sheet.append(['001', index + 1])
        wb.create_sheet('Empty')
        wb.save(path)
        wb.close()
        self.manifest([path])
        result = self.run_script()
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        report = self.report()
        self.assertEqual([d['sheet'] for d in report['datasets']], ['训练', '验证'])
        self.assertEqual(self.imported_rows(report['datasets'][1])[1], ['001', '2'])
        self.assertTrue(any(d.get('sheet') == 'Empty' for d in report['skipped_files']))
        result = self.run_script('build_data_visualization_plan.py')
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        plan = json.loads((self.root / 'paper_output/plan/data_plan.json').read_text(encoding='utf-8'))
        for item in report['datasets']:
            self.assertIn(item['cleaned_output'], json.dumps(plan))

    def test_formula_and_merged_cells_need_explicit_reader(self):
        for kind in ('formula', 'merged'):
            with self.subTest(kind=kind):
                path = self.root / 'problem_files/complex.xlsx'
                wb = Workbook()
                wb.active.append(['x', 'y'])
                wb.active.append([1, '=A2*2' if kind == 'formula' else 2])
                if kind == 'merged':
                    wb.active.merge_cells('A1:B1')
                wb.save(path)
                wb.close()
                self.manifest([path])
                self.assertNotEqual(self.run_script().returncode, 0)
                self.assertEqual(self.report()['status'], 'FAIL')

    def test_same_stem_inputs_do_not_overwrite(self):
        self.manifest([self.source('one/data.csv'), self.source('two/data.csv', 'id,value\n999,2\n')])
        self.assertEqual(self.run_script().returncode, 0)
        outputs = [d['cleaned_output'] for d in self.report()['datasets']]
        self.assertEqual(len(set(outputs)), 2)

    def test_missing_or_invalid_manifest_never_falls_back(self):
        self.source()
        self.assertNotEqual(self.run_script().returncode, 0)
        manifest = self.root / 'paper_output/input_manifest.json'
        for text in ('{broken', '{}', '{"entries":[null]}'):
            with self.subTest(text=text):
                manifest.write_text(text, encoding='utf-8')
                self.assertNotEqual(self.run_script().returncode, 0)
                self.assertNotEqual(self.run_script('robust_loader.py').returncode, 0)

    def test_changed_input_rejected_before_import(self):
        path = self.source()
        self.manifest([path])
        path.write_text('id,value\n001,99\n', encoding='utf-8')
        self.assertNotEqual(self.run_script().returncode, 0)
        self.assertIn('changed', self.report()['errors'][0])

    def test_paths_cannot_escape_project(self):
        manifest = self.manifest([self.source()])
        for value in ('../outside.csv', '/tmp/data.csv', 'C:/outside.csv', 'C:outside.csv', 'a\\..\\outside.csv'):
            with self.subTest(value=value):
                data = json.loads(manifest.read_text(encoding='utf-8'))
                data['entries'][0]['path'] = value
                manifest.write_text(json.dumps(data), encoding='utf-8')
                self.assertNotEqual(self.run_script().returncode, 0)

    def test_chinese_and_windows_separators(self):
        manifest = self.manifest([self.source('数据/观测.csv')])
        data = json.loads(manifest.read_text(encoding='utf-8'))
        data['entries'][0]['path'] = data['entries'][0]['path'].replace('/', '\\')
        manifest.write_text(json.dumps(data), encoding='utf-8')
        self.assertEqual(self.run_script().returncode, 0)
        self.assertEqual(self.report()['datasets'][0]['path'], 'problem_files/数据/观测.csv')

    def test_failed_rerun_invalidates_plan_and_plotting(self):
        path = self.source()
        self.manifest([path])
        self.assertEqual(self.run_script().returncode, 0)
        old = self.root / self.report()['datasets'][0]['cleaned_output']
        path.write_text('x,y\n1,2,3\n', encoding='utf-8')
        self.manifest([path])
        self.assertNotEqual(self.run_script().returncode, 0)
        self.assertTrue(old.exists(), 'previous output is preserved, not reused')
        self.assertNotEqual(self.run_script('build_data_visualization_plan.py').returncode, 0)
        self.assertNotEqual(self.run_script('visualize_data.py').returncode, 0)

    def test_modified_import_invalidates_plan(self):
        self.manifest([self.source()])
        self.assertEqual(self.run_script().returncode, 0)
        output = self.root / self.report()['datasets'][0]['cleaned_output']
        output.write_text('altered', encoding='utf-8')
        self.assertNotEqual(self.run_script('build_data_visualization_plan.py').returncode, 0)

    def test_pipeline_stops_before_plotting_on_bad_data(self):
        self.manifest([self.source(text='x,y\n1,2,3\n')])
        result = self.run_script('run_pipeline.py')
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse((self.root / 'paper_output/figure_index.json').exists())

    def test_s3_and_s6_bind_cleaning_provenance(self):
        self.manifest([self.source()])
        self.assertEqual(self.run_script().returncode, 0)
        output = self.report()['datasets'][0]['cleaned_output']
        modules = []
        for skill, name in (('paper-workflow-orchestrator', 'workflow_guard'),
                            ('quality-assurance-auditor', 'evidence_gate')):
            file = SCRIPTS.parents[1] / skill / 'scripts' / f'{name}.py'
            spec = importlib.util.spec_from_file_location(f'data_test_{name}', file)
            module = importlib.util.module_from_spec(spec)
            with patch('pathlib.Path.cwd', return_value=self.root):
                spec.loader.exec_module(module)
            modules.append(module)
        guard, gate = modules
        hashes = gate.evaluate()['input_hashes']
        self.assertIn(output, hashes)
        self.assertIn('paper_output/data_cleaned/cleaning_report.json', hashes)
        self.assertIn('problem_files/data.csv', hashes)
        (self.root / output).write_text('changed', encoding='utf-8')
        self.assertTrue(any('Data import is stale or failed' in f for f in guard.check_s3()['failures']))
        self.assertTrue(any('Data import is stale or failed' in f for f in gate.evaluate()['failures']))

    def test_project_root_from_another_working_directory(self):
        self.manifest([self.source()])
        result = subprocess.run([sys.executable, str(SCRIPTS / 'clean_data.py'), '--project-root', str(self.root)],
                                cwd=REPO, capture_output=True, encoding='utf-8', timeout=45)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(self.report()['status'], 'PASS')


if __name__ == '__main__':
    unittest.main()
