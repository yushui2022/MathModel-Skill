"""Loss-preserving table import, not automatic statistical preprocessing."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

import pandas as pd

from data_inputs import modeling_inputs, project_path, sha256_file


def find_data_files(project_root: Path | None = None) -> list[Path]:
    return modeling_inputs((project_root or Path.cwd()).resolve())[0]


def table_frame(rows: list[list], label: str) -> pd.DataFrame:
    if not rows:
        return pd.DataFrame()
    header = [str(value) if value is not None else '' for value in rows[0]]
    if any(not name.strip() for name in header) or len(set(header)) != len(header):
        raise ValueError(f'{label}: empty/duplicate header; explicitly resolve the header layout')
    for index, row in enumerate(rows[1:], 2):
        if len(row) != len(header):
            raise ValueError(f'{label}: record {index} has {len(row)} fields, expected {len(header)}; no rows skipped')
    return pd.DataFrame(rows[1:], columns=header, dtype=object).fillna('')


def read_tables(path: Path) -> list[tuple[str | None, pd.DataFrame, dict]]:
    suffix = path.suffix.lower()
    if suffix in {'.csv', '.tsv', '.txt'}:
        for encoding in ('utf-8-sig', 'gb18030'):
            try:
                with path.open('r', encoding=encoding, newline='') as handle:
                    if suffix == '.txt':
                        sample = handle.read(8192)
                        delimiter = csv.Sniffer().sniff(sample, delimiters=',\t;').delimiter
                        handle.seek(0)
                    elif suffix == '.tsv':
                        delimiter = '\t'
                    else:
                        header_sample = handle.readline()
                        handle.seek(0)
                        try:
                            delimiter = csv.Sniffer().sniff(header_sample, delimiters=',\t;').delimiter
                        except csv.Error:
                            delimiter = ','
                    rows = list(csv.reader(handle, delimiter=delimiter, strict=True))
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError(f'{path.name}: unsupported text encoding')
        # A blank CSV record remains an empty observation, never a deleted row.
        if rows:
            rows = [rows[0], *[row if row else [''] * len(rows[0]) for row in rows[1:]]]
        return [(None, table_frame(rows, path.name), {'encoding': encoding, 'delimiter': delimiter})]
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            tables = []
            for sheet in workbook.worksheets:
                if sheet.merged_cells.ranges:
                    raise ValueError(f'{path.name}/{sheet.title}: merged cells require an explicit extraction script')
                rows = []
                for row in sheet.iter_rows():
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError(f'{path.name}/{sheet.title}: formulas require verified evaluated values')
                    rows.append([cell.value if cell.value is not None else '' for cell in row])
                if not any(any(value != '' for value in row) for row in rows):
                    rows = []
                tables.append((sheet.title, table_frame(rows, f'{path.name}/{sheet.title}'), {'header_row': 1}))
            return tables
        finally:
            workbook.close()
    if suffix == '.xls':
        with pd.ExcelFile(path) as workbook:
            return [(name, table_frame(workbook.parse(name, header=None, dtype=object,
                     keep_default_na=False).values.tolist(), f'{path.name}/{name}'), {'header_row': 1})
                    for name in workbook.sheet_names]
    raise ValueError(f'{path.name}: use a task-specific reader for {suffix}; no automatic conversion')


def clean_dataset(path: Path, root: Path, destination: Path) -> tuple[list[dict], list[dict]]:
    source = path.relative_to(root).as_posix()
    records, skipped = [], []
    # Parse every sheet before publishing any sheet from this workbook.
    tables = read_tables(path)
    for index, (sheet, frame, parsing) in enumerate(tables, 1):
        if len(frame.columns) == 0:
            skipped.append({'path': source, 'sheet': sheet, 'reason': 'empty table'})
            continue
        identity = hashlib.sha256(source.encode('utf-8')).hexdigest()[:12]
        output = destination / f'{path.stem}-{identity}-{index}_cleaned.csv'
        frame.to_csv(output, index=False, encoding='utf-8-sig')
        records.append({
            'path': source, 'sheet': sheet, 'source_sha256': sha256_file(path),
            'cleaned_output': output.relative_to(root).as_posix(), 'output_sha256': sha256_file(output),
            'rows_before': len(frame), 'rows_after': len(frame), 'columns': list(frame.columns),
            'empty_cells': {str(col): int((frame[col] == '').sum()) for col in frame.columns},
            'duplicate_rows_preserved': int(frame.duplicated().sum()),
            'parsing': parsing, 'transformations': [],
            'semantic_validation': 'REQUIRED: confirm units, identifiers, missing-value meaning and split policy',
        })
    return records, skipped


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8', errors='replace')
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--project-root', type=Path, default=Path.cwd())
    args = parser.parse_args()
    root = args.project_root.resolve()
    try:
        directory = project_path(root, 'paper_output/data_cleaned')
        report_path = project_path(root, 'paper_output/data_cleaned/cleaning_report.json')
        directory.mkdir(parents=True, exist_ok=True)
        report = {'schema_version': '1.0', 'generated_at': datetime.now(timezone.utc).isoformat(),
                  'generated_by': 'data-cleaning-and-visualization/scripts/clean_data.py',
                  'status': 'RUNNING', 'input_hashes': {}, 'datasets': [], 'skipped_files': [], 'errors': [],
                  'acceptance_scope': 'LOSS_PRESERVING_IMPORT_ONLY'}
        report_path.write_text(json.dumps(report, indent=2), encoding='utf-8')
        try:
            files, report['skipped_files'] = modeling_inputs(root)
            manifest = project_path(root, 'paper_output/input_manifest.json')
            report['input_hashes'] = {'paper_output/input_manifest.json': sha256_file(manifest),
                                     **{p.relative_to(root).as_posix(): sha256_file(p) for p in files}}
            destination = project_path(root, f'paper_output/data_cleaned/runs/{uuid4().hex}')
            destination.mkdir(parents=True)
            for path in files:
                records, skipped = clean_dataset(path, root, destination)
                report['datasets'].extend(records)
                report['skipped_files'].extend(skipped)
            for name, expected in report['input_hashes'].items():
                if sha256_file(project_path(root, name)) != expected:
                    raise ValueError(f'Input changed during import: {name}')
            report['status'] = 'PASS'
        except Exception as exc:
            report['status'] = 'FAIL'
            report['errors'].append(f'{type(exc).__name__}: {exc}')
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(f'[{report["status"]}] {report_path}')
        for error in report['errors']:
            print(error)
        return 0 if report['status'] == 'PASS' else 1
    except (ValueError, OSError) as exc:
        print(f'[FAIL] {exc}')
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
