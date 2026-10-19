"""Fail-closed status and generated-artifact checks; no model calls."""
from __future__ import annotations

import csv
from pathlib import Path

ACCEPTED_STATUSES = frozenset({"computed", "verified", "validated", "pass"})
IMAGE_SUFFIXES = frozenset({".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff"})
TABLE_SUFFIXES = frozenset({".csv", ".tsv", ".xlsx"})


def status_errors(item: dict, label: str, *, required: bool = False) -> list[str]:
    fields = [key for key in ("status", "evidence_status") if key in item]
    errors = []
    if required and not fields:
        errors.append(f"{label}: missing evidence status")
    for key in fields:
        value = item[key]
        if not isinstance(value, str) or value.strip().casefold() not in ACCEPTED_STATUSES:
            errors.append(f"{label}: unacceptable {key}: {value!r}")
    if item.get("evidence_role") == "exploratory":
        errors.append(f"{label}: exploratory output is not formal evidence")
    return errors


def _check_rows(rows) -> None:
    header = next(rows, None)
    if not header:
        raise ValueError("table has no header")
    names = [str(value).strip() if value is not None else "" for value in header]
    if any(not name for name in names) or len(set(names)) != len(names):
        raise ValueError("table header is empty or duplicated")
    count = 0
    for row in rows:
        if len(row) != len(header):
            raise ValueError("table row width differs from header")
        if any(value is not None and str(value).strip() for value in row):
            count += 1
    if not count:
        raise ValueError("table has no data rows")


def artifact_errors(path: Path, *, kind: str | None = None) -> list[str]:
    """Validate generated raster images and data tables, not arbitrary raw inputs."""
    try:
        if not path.is_file() or path.stat().st_size == 0:
            raise ValueError("artifact is missing or empty")
        suffix = path.suffix.casefold()
        if kind == "figure" and suffix not in IMAGE_SUFFIXES:
            raise ValueError("formal figures must be readable raster images")
        if kind == "table" and suffix not in TABLE_SUFFIXES:
            raise ValueError("formal tables must be CSV, TSV or XLSX")
        if suffix in IMAGE_SUFFIXES:
            from PIL import Image
            with Image.open(path) as image:
                image.verify()
            with Image.open(path) as image:
                image.load()
        elif suffix in {".csv", ".tsv"}:
            with path.open(encoding="utf-8-sig", newline="") as handle:
                _check_rows(iter(csv.reader(handle, delimiter="\t" if suffix == ".tsv" else ",", strict=True)))
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in workbook:
                    _check_rows(iter(sheet.iter_rows(values_only=True)))
            finally:
                workbook.close()
    except Exception as exc:
        return [f"invalid artifact {path.name}: {exc}"]
    return []
